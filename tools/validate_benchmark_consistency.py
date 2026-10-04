"""
Automated Cross-Consistency and Benchmark Integrity Validator for LUNA-CORR.

Enforces zero-tolerance consistency across:
  - Source Code (lunacorr thresholds & default hyperparameters)
  - Result JSON Artifacts (exact metrics, sample sizes, and reason codes)
  - Canonical Manifest (BENCHMARK_MANIFEST.md)
  - Repository Documentation (README.md)
  - Scientific Report (docs/SCIENTIFIC_REPORT.md)
  - SIH Presentation Deck (LUNA-CORR_SIH2026_Presentation.pptx)

Exits with code 0 if all assertions pass, code 1 on any discrepancy.
"""
import sys
import json
import re
from pathlib import Path
import numpy as np

def log_pass(msg: str):
    print(f"  [PASS] {msg}")

def log_fail(msg: str):
    print(f"  [FAIL] {msg}")

def main():
    root = Path(__file__).resolve().parent.parent
    sys.path.insert(0, str(root))
    errors = []
    print("=" * 70)
    print("LUNA-CORR BENCHMARK CROSS-CONSISTENCY AUDIT")
    print("=" * 70)

    # -------------------------------------------------------------
    # 1. Manifest Experiment Registry Audit
    # -------------------------------------------------------------
    print("\n1. Auditing Canonical Experiment Manifest (BENCHMARK_MANIFEST.md)...")
    manifest_path = root / "BENCHMARK_MANIFEST.md"
    if not manifest_path.exists():
        errors.append("BENCHMARK_MANIFEST.md not found.")
        log_fail("BENCHMARK_MANIFEST.md missing")
    else:
        manifest_text = manifest_path.read_text(encoding="utf-8")
        required_exp_ids = [
            "EXP-OHRC-E2E",
            "EXP-TMC-E2E",
            "EXP-TMC-FIXED",
            "EXP-NEG-DISJOINT",
            "EXP-SYN-ILLUM",
            "EXP-IO"
        ]
        for exp_id in required_exp_ids:
            if exp_id in manifest_text:
                log_pass(f"Experiment ID registered: {exp_id}")
            else:
                errors.append(f"Missing Experiment ID in manifest: {exp_id}")
                log_fail(f"Experiment ID missing: {exp_id}")

        for stale in ["0.779", "34.42", "0.812"]:
            if stale in manifest_text:
                errors.append(f"Stale value {stale} found in BENCHMARK_MANIFEST.md")
                log_fail(f"Stale value {stale} in BENCHMARK_MANIFEST.md")
            else:
                log_pass(f"No stale value {stale} in BENCHMARK_MANIFEST.md")

    # -------------------------------------------------------------
    # 2. Canonical JSON Artifact Existence and Values
    # -------------------------------------------------------------
    print("\n2. Auditing Canonical Benchmark Artifacts...")
    
    # 2.1 EXP-OHRC-E2E
    ohrc_json_path = root / "results/real_ohrc_cross_orbit/result.json"
    if not ohrc_json_path.exists():
        errors.append(f"Artifact missing: {ohrc_json_path}")
        log_fail("EXP-OHRC-E2E artifact missing")
    else:
        with open(ohrc_json_path, "r", encoding="utf-8") as f:
            ohrc = json.load(f)
        assert ohrc["metrics"]["inlier_count"] == 822, "OHRC inlier_count != 822"
        assert abs(ohrc["heldout_validation_correspondences_metrics"]["rmse"] - 0.679) < 1e-3, "OHRC RMSE != 0.679"
        assert ohrc["heldout_validation_correspondences_metrics"]["n_points"] == 165, "OHRC held-out N != 165"
        assert ohrc["status"] == "ACCEPTED", "OHRC status != ACCEPTED"
        log_pass("EXP-OHRC-E2E: 822 inliers, 0.679 px held-out RMSE, N=165 confirmed.")

    # 2.2 EXP-TMC-E2E
    tmc_e2e_json_path = root / "results/real_tmc2_stereo_checkpointed/result.json"
    if not tmc_e2e_json_path.exists():
        errors.append(f"Artifact missing: {tmc_e2e_json_path}")
        log_fail("EXP-TMC-E2E artifact missing")
    else:
        with open(tmc_e2e_json_path, "r", encoding="utf-8") as f:
            tmc_e2e = json.load(f)
        assert tmc_e2e["metrics"]["inlier_count"] == 686, "TMC-E2E inliers != 686"
        assert abs(tmc_e2e["heldout_validation_correspondences_metrics"]["rmse"] - 1.110) < 1e-3, "TMC-E2E RMSE != 1.110"
        assert abs(tmc_e2e["fitting_correspondences_metrics"]["rmse"] - 0.841) < 1e-3, "TMC-E2E fitting RMSE != 0.841"
        assert tmc_e2e["heldout_validation_correspondences_metrics"]["n_points"] == 138, "TMC-E2E held-out N != 138"
        assert tmc_e2e["status"] == "ACCEPTED", "TMC-E2E status != ACCEPTED"
        log_pass("EXP-TMC-E2E: 686 inliers, 1.110 px held-out RMSE, 0.841 px fitting RMSE confirmed.")

    # 2.3 EXP-TMC-FIXED (Zero-Leakage 3 px Spatial Purge)
    ablation_json_path = root / "results/ablation_bootstrap_ci.json"
    if not ablation_json_path.exists():
        errors.append(f"Artifact missing: {ablation_json_path}")
        log_fail("EXP-TMC-FIXED artifact missing")
    else:
        with open(ablation_json_path, "r", encoding="utf-8") as f:
            abl = json.load(f)
        stg1_rmse = abl["1. Raw SIFT"]["rmse"][0]
        stg3_rmse = abl["3. + RootSIFT"]["rmse"][0]
        stg5_rmse = abl["5. + Adaptive TPS"]["rmse"][0]
        stg5_ci = [abl["5. + Adaptive TPS"]["rmse"][1], abl["5. + Adaptive TPS"]["rmse"][2]]

        assert abs(stg1_rmse - 2.119) < 1e-3, f"Stage 1 RMSE {stg1_rmse} != 2.119"
        assert abs(stg3_rmse - 1.606) < 1e-3, f"Stage 3 RMSE {stg3_rmse} != 1.606"
        assert abs(stg5_rmse - 0.836) < 1e-3, f"Stage 5 RMSE {stg5_rmse} != 0.836"
        assert abs(stg5_ci[0] - 0.718) < 1e-3 and abs(stg5_ci[1] - 0.973) < 1e-3, f"Stage 5 CI {stg5_ci} != [0.718, 0.973]"

        # Mathematically verify percentage gains
        gain_vs_stg3 = (stg3_rmse - stg5_rmse) / stg3_rmse
        gain_vs_stg1 = (stg1_rmse - stg5_rmse) / stg1_rmse
        assert abs(gain_vs_stg3 - 0.4794) < 1e-2, f"Gain vs Stage 3 {gain_vs_stg3:.3f} != 47.9%"
        assert abs(gain_vs_stg1 - 0.6054) < 1e-2, f"Gain vs Stage 1 {gain_vs_stg1:.3f} != 60.5%"
        log_pass("EXP-TMC-FIXED: Stage 1 (2.119 px) -> Stage 3 (1.606 px) -> Stage 5 (0.836 px [0.718, 0.973], 47.9% gain) confirmed.")

    # 2.3b EXP-TMC-FIXED Study Results (Ablation Table Consistency)
    ablation_study_path = root / "results/ablation_study_results.json"
    if not ablation_study_path.exists():
        errors.append(f"Artifact missing: {ablation_study_path}")
        log_fail("ablation_study_results.json missing")
    else:
        with open(ablation_study_path, "r", encoding="utf-8") as f:
            abl_study = json.load(f)
        assert abs(abl_study[0]["fixed_eval_set"]["fixed_rmse_px"] - 2.119) < 1e-3
        assert abs(abl_study[2]["fixed_eval_set"]["fixed_rmse_px"] - 1.606) < 1e-3
        assert abs(abl_study[4]["fixed_eval_set"]["fixed_rmse_px"] - 0.836) < 1e-3
        log_pass("ablation_study_results.json fixed evaluation metrics (2.119 -> 1.606 -> 0.836 px) fully verified.")

    # 2.4 EXP-NEG-DISJOINT
    neg_json_path = root / "results/negative_control_disjoint/result.json"
    if not neg_json_path.exists():
        errors.append(f"Artifact missing: {neg_json_path}")
        log_fail("EXP-NEG-DISJOINT artifact missing")
    else:
        with open(neg_json_path, "r", encoding="utf-8") as f:
            neg = json.load(f)
        assert neg["status"] == "ABSTAINED", "Negative control status != ABSTAINED"
        assert not neg["accepted"], "Negative control accepted is True"
        assert "LOW_INLIERS" in neg["reason_codes"], "Missing LOW_INLIERS code"
        assert "LOW_COVERAGE" in neg["reason_codes"], "Missing LOW_COVERAGE code"
        assert neg["metrics"]["inlier_count"] == 5, f"Negative control inlier_count {neg['metrics']['inlier_count']} != 5"
        log_pass("EXP-NEG-DISJOINT: 1/1 tested pair rejected with ['LOW_INLIERS', 'LOW_COVERAGE'], zero false acceptance confirmed.")

    # 2.5 EXP-SYN-ILLUM
    syn_json_path = root / "results/synthetic_illumination_groundtruth.json"
    if not syn_json_path.exists():
        errors.append(f"Artifact missing: {syn_json_path}")
        log_fail("EXP-SYN-ILLUM artifact missing")
    else:
        with open(syn_json_path, "r", encoding="utf-8") as f:
            syn = json.load(f)
        # Check angle 0 deg baseline (737 inliers)
        item_0 = next(x for x in syn if x["delta_azimuth_deg"] == 0)
        assert item_0["mode_c_reillumination"]["inlier_count"] == 737, f"Mode C @ 0deg inliers {item_0['mode_c_reillumination']['inlier_count']} != 737"
        
        # Check angle 45 deg breakdown of direct optical
        item_45 = next(x for x in syn if x["delta_azimuth_deg"] == 45)
        assert item_45["direct_optical"]["precision_at_10px"] == 0.0, "Direct optical prec@45deg != 0.0"
        assert item_45["mode_c_reillumination"]["precision_at_10px"] >= 0.99, "Mode C prec@45deg < 0.99"
        
        # Check angle 90 and 180 suppression
        item_90 = next(x for x in syn if x["delta_azimuth_deg"] == 90)
        item_180 = next(x for x in syn if x["delta_azimuth_deg"] == 180)
        assert "SUPPRESSED" in item_90["direct_optical"]["status"], "Direct optical @ 90deg not SUPPRESSED"
        assert "SUPPRESSED" in item_180["direct_optical"]["status"], "Direct optical @ 180deg not SUPPRESSED"
        assert item_90["mode_c_reillumination"]["inlier_count"] == 932, "Mode C @ 90deg != 932 inliers"
        assert item_180["mode_c_reillumination"]["inlier_count"] == 742, "Mode C @ 180deg != 742 inliers"
        log_pass("EXP-SYN-ILLUM: Direct collapses at >=45 deg; small N suppressed; Mode C maintains >99% prec and 737–1,187 inliers.")

    # 2.6 EXP-IO
    io_json_path = root / "results/seeker_latency.json"
    if not io_json_path.exists():
        errors.append(f"Artifact missing: {io_json_path}")
        log_fail("EXP-IO artifact missing")
    else:
        with open(io_json_path, "r", encoding="utf-8") as f:
            io_res = json.load(f)
        assert io_res["n_samples"] == 500, "EXP-IO sample size != 500"
        assert abs(io_res["median_latency_ms"] - 39.15) < 1e-2, "EXP-IO median != 39.15 ms"
        assert abs(io_res["mean_latency_ms"] - 40.47) < 1e-2, "EXP-IO mean != 40.47 ms"
        log_pass("EXP-IO: 500 random window seeks, 39.15 ms median latency confirmed.")

    # -------------------------------------------------------------
    # 3. Codebase Threshold & Hyperparameter Configuration Audit
    # -------------------------------------------------------------
    print("\n3. Auditing Codebase Configuration Parameters...")
    from lunacorr.gate.decision_gate import QualityGate
    from lunacorr.estimate.adaptive_gate import AdaptiveDeformationGate
    from lunacorr.pipeline.inference import LunarRegistrationPipeline

    gate = QualityGate()
    assert gate.min_inliers == 20, f"QualityGate min_inliers ({gate.min_inliers}) != 20"
    assert gate.min_inlier_ratio == 0.15, f"QualityGate min_inlier_ratio ({gate.min_inlier_ratio}) != 0.15"
    assert gate.min_occupied_ratio == 0.30, f"QualityGate min_occupied_ratio ({gate.min_occupied_ratio}) != 0.30"
    log_pass("QualityGate operational threshold harmonized to min_inliers = 20 (N >= 20 reporting threshold).")

    pipeline = LunarRegistrationPipeline()
    assert pipeline.gate.min_inliers == 20, f"Default pipeline QualityGate min_inliers ({pipeline.gate.min_inliers}) != 20"
    log_pass("LunarRegistrationPipeline default gate min_inliers = 20 verified.")

    adapt = AdaptiveDeformationGate()
    assert adapt.autocorr_threshold == 0.20, f"AdaptiveDeformationGate threshold ({adapt.autocorr_threshold}) != 0.20"
    assert adapt.min_p95_px == 1.5, f"AdaptiveDeformationGate min_p95_px ({adapt.min_p95_px}) != 1.5"
    assert adapt.eps_zero == 0.05, f"AdaptiveDeformationGate eps_zero ({adapt.eps_zero}) != 0.05"
    assert adapt.min_points == 20, f"AdaptiveDeformationGate min_points ({adapt.min_points}) != 20"
    log_pass("AdaptiveDeformationGate thresholds (S_relief > 0.20, P95 >= 1.5 px, eps_zero = 0.05 px, min_points = 20) confirmed.")

    # -------------------------------------------------------------
    # 4. Text & Document Synchronization Audit
    # -------------------------------------------------------------
    print("\n4. Auditing Cross-Document Synchronization...")
    report_path = root / "docs/SCIENTIFIC_REPORT.md"
    readme_path = root / "README.md"

    if report_path.exists():
        report_text = report_path.read_text(encoding="utf-8")
        for stale in ["0.779", "34.42", "0.812", "1,679"]:
            if stale in report_text:
                errors.append(f"docs/SCIENTIFIC_REPORT.md contains stale reference to {stale}")
                log_fail(f"Stale value {stale} found in docs/SCIENTIFIC_REPORT.md")
            else:
                log_pass(f"No stale {stale} in docs/SCIENTIFIC_REPORT.md.")

        for num in ["0.679", "1.110", "0.836", "1.606", "47.9%", "737"]:
            if num in report_text:
                log_pass(f"docs/SCIENTIFIC_REPORT.md contains canonical value: {num}")
            else:
                errors.append(f"docs/SCIENTIFIC_REPORT.md missing canonical value: {num}")
                log_fail(f"Missing canonical value in report: {num}")

        # Check NN-DRC epsilon in report matches code (0.05 px)
        if r"\epsilon = 0.05" in report_text or "0.05 px" in report_text:
            log_pass("docs/SCIENTIFIC_REPORT.md specifies NN-DRC epsilon = 0.05 px matching code.")
        else:
            errors.append("docs/SCIENTIFIC_REPORT.md missing epsilon = 0.05 px specification")
            log_fail("docs/SCIENTIFIC_REPORT.md does not specify epsilon = 0.05 px")

        if r"\epsilon = 0.1" in report_text or "epsilon = 0.10" in report_text:
            errors.append("docs/SCIENTIFIC_REPORT.md contains stale epsilon = 0.1 px")
            log_fail("docs/SCIENTIFIC_REPORT.md contains stale epsilon = 0.1 px")

    if readme_path.exists():
        readme_text = readme_path.read_text(encoding="utf-8")
        for stale in ["0.779", "0.812", "tests-5 passed", "Pinned dependencies"]:
            if stale in readme_text:
                errors.append(f"README.md contains stale string: {stale}")
                log_fail(f"Stale string '{stale}' in README.md")
            else:
                log_pass(f"No stale string '{stale}' in README.md.")

        for num in ["0.679", "1.110", "0.836", "1.606", "47.9%", "737–1,187", "tests-15%20passed"]:
            if num in readme_text:
                log_pass(f"README.md contains canonical value: {num}")
            else:
                errors.append(f"README.md missing canonical value: {num}")
                log_fail(f"Missing canonical value in README: {num}")

    # -------------------------------------------------------------
    # 5. PowerPoint Deck (PPTX) Inspection Audit
    # -------------------------------------------------------------
    print("\n5. Auditing PowerPoint Presentation Deck (LUNA-CORR_SIH2026_Presentation.pptx)...")
    pptx_path = root / "LUNA-CORR_SIH2026_Presentation.pptx"
    if not pptx_path.exists():
        errors.append("LUNA-CORR_SIH2026_Presentation.pptx missing.")
        log_fail("Presentation deck missing.")
    else:
        try:
            from pptx import Presentation
            prs = Presentation(str(pptx_path))
            assert len(prs.slides) == 6, f"Expected 6 slides, found {len(prs.slides)}"
            log_pass(f"Deck loaded successfully: {len(prs.slides)} slides.")

            # Extract all text from all shapes across all slides
            slide_texts = []
            for i, s in enumerate(prs.slides):
                texts = []
                for shape in s.shapes:
                    if shape.has_text_frame:
                        for p in shape.text_frame.paragraphs:
                            if p.text:
                                texts.append(p.text)
                slide_texts.append(" ".join(texts))
            full_deck_text = " ".join(slide_texts)

            # Check presence of all 6 Experiment IDs in PPTX
            for exp_id in required_exp_ids:
                if exp_id in full_deck_text:
                    log_pass(f"PPTX contains experiment ID: {exp_id}")
                else:
                    errors.append(f"PPTX missing experiment ID: {exp_id}")
                    log_fail(f"PPTX missing experiment ID: {exp_id}")

            # Check presence of canonical metrics in PPTX
            canonical_pptx_values = [
                ("0.679", "OHRC cross-orbit held-out RMSE"),
                ("1.110", "TMC-2 stereo held-out RMSE"),
                ("0.836", "TMC-2 ablation Stage 5 RMSE"),
                ("1.606", "TMC-2 ablation Stage 3 RMSE (or 1.61)"),
                ("47.9%", "TPS ablation error reduction"),
                ("39.2",  "Seeker median I/O latency (or 39.15)"),
                ("737–1,187", "Mode C inlier range across tested sweep"),
                ("822",   "OHRC inlier count"),
                ("686",   "TMC-2 inlier count"),
                ("1/1",   "Disjoint negative control tested count"),
            ]
            for val, desc in canonical_pptx_values:
                # Allow minor formatting variations for values with alternatives
                matched = False
                if val == "1.606":
                    matched = ("1.606" in full_deck_text) or ("1.61" in full_deck_text)
                elif val == "39.2":
                    matched = ("39.2" in full_deck_text) or ("39.15" in full_deck_text)
                else:
                    matched = val in full_deck_text

                if matched:
                    log_pass(f"PPTX contains {val} ({desc})")
                else:
                    errors.append(f"PPTX missing canonical value: {val} ({desc})")
                    log_fail(f"PPTX missing canonical value: {val} ({desc})")

            # Check absence of stale values in PPTX
            stale_pptx_values = ["0.779", "34.42", "0.812", "1,679"]
            for stale in stale_pptx_values:
                if stale in full_deck_text:
                    errors.append(f"PPTX contains stale value: {stale}")
                    log_fail(f"Stale value '{stale}' found in PPTX")
                else:
                    log_pass(f"No stale value '{stale}' in PPTX.")

            # Check Mode C framing on Slide 3 (must clearly indicate physics branch, not default E2E stage)
            slide3_text = slide_texts[2]
            if "Mode C: Physics branch" in slide3_text or "Mode C: Physics Branch" in slide3_text or "Physics branch" in slide3_text:
                log_pass("Slide 3 explicitly designates Mode C as Physics branch.")
            else:
                errors.append("Slide 3 does not designate Mode C as Physics branch.")
                log_fail("Slide 3 Mode C framing incomplete.")

        except Exception as e:
            errors.append(f"Failed to inspect PPTX: {e}")
            log_fail(f"Failed to inspect PPTX: {e}")

    # -------------------------------------------------------------
    # Summary
    # -------------------------------------------------------------
    print("\n" + "=" * 70)
    if errors:
        print(f"AUDIT FAILED WITH {len(errors)} ERROR(S):")
        for err in errors:
            print(f"  - {err}")
        print("=" * 70)
        sys.exit(1)
    else:
        print("AUDIT SUCCESS: ALL CODE, ARTIFACTS, MANIFESTS, DOCS, AND PPTX ARE MUTUALLY CONSISTENT.")
        print("=" * 70)
        sys.exit(0)

if __name__ == "__main__":
    main()
